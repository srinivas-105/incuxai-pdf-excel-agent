import io
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def make_workbook(tables, warnings, page_log):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        used_names = set()
        if not tables:
            pd.DataFrame([{"message": "No tables were detected."}]).to_excel(writer, sheet_name="Summary", index=False)
        for i, table in enumerate(tables, start=1):
            base = f"Table_{i}"
            name = base[:31]
            suffix = 1
            while name in used_names:
                suffix += 1
                name = f"{base[:26]}_{suffix}"[:31]
            used_names.add(name)
            df = pd.DataFrame(table["data"], columns=table["columns"])
            df.to_excel(writer, sheet_name=name, index=False)
        pd.DataFrame(warnings or [{"severity": "info", "issue": "No structural warnings detected."}]).to_excel(writer, sheet_name="Validation", index=False)
        pd.DataFrame(page_log or [{"note": "No pages processed."}]).to_excel(writer, sheet_name="Processing_Log", index=False)

        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="174A7E")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for column_cells in ws.columns:
                max_len = 0
                col_letter = get_column_letter(column_cells[0].column)
                for cell in column_cells[:200]:
                    max_len = max(max_len, len(str(cell.value or "")))
                ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 40)
    output.seek(0)
    return output.getvalue()
